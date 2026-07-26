"""CLIN / price-schedule structure ingestion (PRD §9.1, FR-12).

Government XLSX pricing templates are detected and preserved as fillable
artifacts — never flattened. When a template can't be safely machine-filled
(macros, merged cells), the fallback is "human fills, system computes."
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from ..models import DocTree
from ..routing import ModelRouter, Tier
from .models import PricingStructure

SYSTEM = """You are a federal pricing analyst. Extract the pricing structure the
government has defined: every CLIN/SLIN or price-schedule line, its description,
contract type per CLIN (FFP, T&M, cost-reimbursable) if determinable, quantities/
units, and period (base vs option years). Note any IGCE hints, historical award
values, or funding ceilings mentioned anywhere in the documents. Do not invent
CLINs — if no explicit pricing structure exists, return an empty list and explain
in notes."""


def parse_structure(router: ModelRouter, doc_tree: DocTree) -> PricingStructure:
    structure = router.structured(
        Tier.FRONTIER,
        system=SYSTEM,
        prompt=f"Extract the pricing structure from this solicitation corpus.\n\n{doc_tree.corpus()}",
        output_type=PricingStructure,
        stage="pricing.structure",
    )
    template = find_government_template(doc_tree)
    if template:
        structure.government_template_file = template
    return structure


def find_government_template(doc_tree: DocTree) -> Optional[str]:
    """Locate a government-provided XLSX pricing template among the attachments."""
    for doc in doc_tree.docs:
        if doc.kind == "xlsx" and doc.fillable_template:
            name_lower = doc.name.lower()
            if any(k in name_lower for k in ("price", "pricing", "cost", "clin", "schedule", "bid")):
                return doc.name
    for doc in doc_tree.docs:
        if doc.kind == "xlsx" and doc.fillable_template:
            return doc.name
    return None


WD_SYSTEM = """Extract the Service Contract Act / Davis-Bacon wage determination
table from this solicitation's attachments, if one is present. For each labor
category listed in the WD: the minimum hourly wage and the health & welfare
hourly benefit. Also capture the WD number and locality. If NO wage
determination is attached, return an empty entries list — do not invent one."""


def extract_wage_determination(router: ModelRouter, doc_tree: DocTree):
    """Parse the attached WD table (PRD §9.4) so the rate engine can enforce
    the legal floors in code. Returns None when no WD is present."""
    from .rates import WageDetermination

    corpus = doc_tree.corpus()
    if "wage determination" not in corpus.lower() and "wage rate" not in corpus.lower():
        return None
    wd = router.structured(
        Tier.FRONTIER,
        system=WD_SYSTEM,
        prompt=f"Extract the wage determination.\n\n{corpus}",
        output_type=WageDetermination,
        stage="pricing.wd",
    )
    return wd if wd.entries else None


def describe_template(path: str | Path) -> dict:
    """List sheets/dimensions of a government XLSX template so a human (or a
    later fill pass) knows what needs filling. Formulas are preserved because
    we never rewrite the file, only copy it into the package."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=False)
    info = {"sheets": []}
    for ws in wb.worksheets:
        info["sheets"].append(
            {"name": ws.title, "max_row": ws.max_row, "max_col": ws.max_column}
        )
    return info
