"""Government XLSX pricing-template filling (FR-12, PRD §9.6).

The rule: fill THEIR file, never generate our own. Approach:
1. Code dumps the template's cell grid (labels + coordinates).
2. The frontier model proposes a cell mapping: which cells take which of our
   computed numbers (a proposal, never trusted blindly).
3. Code applies the writes to a COPY of the template with openpyxl,
   preserving formulas, formatting, and everything untouched.
4. Guardrails: macro-enabled workbooks (.xlsm) and formula-target cells are
   never overwritten — those fall back to "human fills, system computes",
   with the proposal recorded so the human knows exactly what goes where.
"""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Optional

from ..routing import ModelRouter, Tier
from .models import PricingModel, TemplateFillProposal

MAX_GRID_CELLS = 4_000

MAPPING_SYSTEM = """You map a contractor's computed pricing numbers into the
government's own Excel pricing template. You get the template's cell grid
(sheet, cell reference, current value) and the computed priced lines/totals.

Rules:
- Propose writes ONLY into cells that are clearly meant for offeror input
  (empty cells next to labels like 'Unit Price', 'Extended Price', 'Total',
  labor-category rate columns, CLIN rows).
- NEVER propose writing over a cell that contains a formula (values starting
  with '=') or a label.
- Numbers as plain digits (no $ or commas). Rates are fully-burdened wrapped
  rates unless the template clearly asks for direct rates.
- If the structure is unclear, ambiguous, or the numbers don't map cleanly,
  return an empty writes list and set unfillable_reason — the fallback is
  'human fills, system computes' and a wrong guess is worse than no guess."""


def propose_fill(
    router: ModelRouter, template_path: Path, pricing: PricingModel
) -> TemplateFillProposal:
    if template_path.suffix.lower() == ".xlsm":
        return TemplateFillProposal(
            unfillable_reason="Macro-enabled workbook (.xlsm) — never machine-filled; "
            "human fills, system computes."
        )
    grid = dump_grid(template_path)
    if grid is None:
        return TemplateFillProposal(unfillable_reason="Template could not be read.")
    lines = "\n".join(
        f"- {l.labor_category} year {l.year}: {l.hours} hrs @ wrapped ${l.wrapped_rate}/hr "
        f"= ${l.extended}"
        for l in pricing.priced_lines
    )
    prompt = f"""Map these computed numbers into the template.

=== COMPUTED PRICING ===
{lines}
TOTAL: ${pricing.total}
ODCs: {[f"{o.description}: {o.estimated_cost or 'QUOTE NEEDED'}" for o in pricing.odcs]}

=== TEMPLATE CELL GRID ===
{grid}"""
    return router.structured(
        Tier.FRONTIER,
        system=MAPPING_SYSTEM,
        prompt=prompt,
        output_type=TemplateFillProposal,
        stage="pricing.template_fill",
    )


def dump_grid(template_path: Path) -> Optional[str]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(str(template_path), data_only=False)
    except Exception:
        return None
    rows: list[str] = []
    count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                rows.append(f"{ws.title}!{cell.coordinate}: {str(cell.value)[:80]}")
                count += 1
                if count >= MAX_GRID_CELLS:
                    rows.append("[grid truncated]")
                    return "\n".join(rows)
    return "\n".join(rows) or "(empty workbook)"


def apply_fill(
    template_path: Path, proposal: TemplateFillProposal, dest_path: Path
) -> tuple[Optional[Path], list[str]]:
    """Apply proposed writes to a COPY of the template. Returns (filled path
    or None, skipped-write notes). Formula cells are never overwritten."""
    import openpyxl

    if proposal.unfillable_reason or not proposal.writes:
        return None, [proposal.unfillable_reason or "no writes proposed"]

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template_path, dest_path)
    skipped: list[str] = []
    try:
        wb = openpyxl.load_workbook(str(dest_path))
        for write in proposal.writes:
            if write.sheet not in wb.sheetnames:
                skipped.append(f"{write.sheet}!{write.cell}: sheet not found")
                continue
            ws = wb[write.sheet]
            try:
                cell = ws[write.cell]
            except Exception:
                skipped.append(f"{write.sheet}!{write.cell}: bad reference")
                continue
            if isinstance(cell.value, str) and cell.value.startswith("="):
                skipped.append(f"{write.sheet}!{write.cell}: formula cell — never overwritten")
                continue
            value = write.value
            try:
                cell.value = float(value) if _is_number(value) else value
            except Exception:
                skipped.append(f"{write.sheet}!{write.cell}: could not set value")
        wb.save(str(dest_path))
    except Exception as exc:
        dest_path.unlink(missing_ok=True)
        return None, [f"fill failed, template untouched: {exc}"]
    return dest_path, skipped


def _is_number(value: str) -> bool:
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False
