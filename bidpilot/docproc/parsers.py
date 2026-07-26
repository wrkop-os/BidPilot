"""Document processing (A2): the parsing ladder.

PDF: native text (pypdf) -> table extraction (pdfplumber, when installed)
-> OCR fallback flag for image-only pages (vision-LLM transcription is the
production path; v1 flags the doc for manual review rather than silently
emitting garbage). DOCX/XLSX parsed natively; XLSX preserved as fillable
artifacts, never flattened. Every doc is scanned for CUI/ITAR markings,
which trigger the halt-and-notify path (NG2, §14.5).
"""

from __future__ import annotations

import re
from pathlib import Path

from ..models import AttachmentRecord, DocSection, DocTree, ParsedDoc, TableData

MAX_CHARS_PER_DOC = 400_000

CUI_PATTERNS = [
    r"\bCUI\b(?!\w)",
    r"CONTROLLED UNCLASSIFIED INFORMATION",
    r"\bITAR\b",
    r"EXPORT[- ]CONTROLLED",
    r"\bFOUO\b",
    r"DISTRIBUTION STATEMENT [B-F]",
]

# UCF Section A–M headers plus common attachment/PWS markers.
SECTION_HEADER_RE = re.compile(
    r"^\s*(SECTION\s+([A-M])\b[^\n]*|PART\s+[IVX]+[^\n]*|"
    r"(ATTACHMENT|EXHIBIT|APPENDIX)\s+\w+[^\n]*|"
    r"(STATEMENT OF WORK|PERFORMANCE WORK STATEMENT)[^\n]*)\s*$",
    re.IGNORECASE | re.MULTILINE,
)


def process_attachments(attachments: list[AttachmentRecord], description_text: str = "") -> DocTree:
    docs: list[ParsedDoc] = []
    if description_text.strip():
        docs.append(
            ParsedDoc(
                name="SAM.gov notice description",
                kind="other",
                full_text=description_text[:MAX_CHARS_PER_DOC],
                sections=detect_sections(description_text),
                cui_markings=scan_cui(description_text),
            )
        )
    for att in attachments:
        if att.restricted or not Path(att.local_path).exists():
            continue
        docs.append(parse_file(att.local_path))
    return DocTree(docs=docs)


def parse_file(path: str | Path) -> ParsedDoc:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix in (".docx", ".dotx"):
        return _parse_docx(path)
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return _parse_xlsx(path)
    if suffix in (".txt", ".md", ".csv", ".json", ".xml", ".htm", ".html"):
        text = path.read_text(encoding="utf-8", errors="replace")[:MAX_CHARS_PER_DOC]
        return ParsedDoc(
            name=path.name, kind="other", full_text=text,
            sections=detect_sections(text), cui_markings=scan_cui(text),
        )
    if suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
        return ParsedDoc(name=path.name, kind="image", ocr_used=True, full_text="")
    return ParsedDoc(name=path.name, kind="other")


def _parse_pdf(path: Path) -> ParsedDoc:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception:
        return ParsedDoc(name=path.name, kind="pdf")

    page_texts: list[str] = []
    empty_pages = 0
    for page in reader.pages:
        text = ""
        try:
            text = page.extract_text() or ""
        except Exception:
            pass
        if not text.strip():
            empty_pages += 1
        page_texts.append(text)

    full_text = "\n\n".join(
        f"[page {i + 1}]\n{t}" for i, t in enumerate(page_texts)
    )[:MAX_CHARS_PER_DOC]

    tables = _extract_pdf_tables(path)
    # Image-only PDF: >60% empty pages means the native ladder failed — OCR needed.
    ocr_needed = len(page_texts) > 0 and empty_pages / len(page_texts) > 0.6

    return ParsedDoc(
        name=path.name,
        kind="pdf",
        page_count=len(page_texts),
        full_text=full_text,
        sections=detect_sections(full_text),
        tables=tables,
        ocr_used=ocr_needed,
        cui_markings=scan_cui(full_text),
    )


def _extract_pdf_tables(path: Path, max_tables: int = 40) -> list[TableData]:
    """Preserve tables as tables (pricing schedules!). pdfplumber is optional;
    without it, tables remain embedded in the text stream."""
    try:
        import pdfplumber
    except ImportError:
        return []
    tables: list[TableData] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                for raw in page.extract_tables() or []:
                    rows = [[(cell or "").strip() for cell in row] for row in raw]
                    if any(any(cell for cell in row) for row in rows):
                        tables.append(TableData(doc_name=path.name, page=page_num, rows=rows))
                        if len(tables) >= max_tables:
                            return tables
    except Exception:
        pass
    return tables


def _parse_docx(path: Path) -> ParsedDoc:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception:
        return ParsedDoc(name=path.name, kind="docx")
    parts = [p.text for p in document.paragraphs]
    tables: list[TableData] = []
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append(TableData(doc_name=path.name, rows=rows))
        parts.extend(" | ".join(r) for r in rows)
    full_text = "\n".join(parts)[:MAX_CHARS_PER_DOC]
    return ParsedDoc(
        name=path.name, kind="docx", full_text=full_text,
        sections=detect_sections(full_text), tables=tables, cui_markings=scan_cui(full_text),
    )


def _parse_xlsx(path: Path) -> ParsedDoc:
    """XLSX = potential government pricing template: read cell text for the
    corpus but mark the file as a fillable artifact — the original is copied
    into the package untouched, never regenerated."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception:
        return ParsedDoc(name=path.name, kind="xlsx", fillable_template=True)
    parts = []
    tables: list[TableData] = []
    for ws in wb.worksheets:
        parts.append(f"[sheet: {ws.title}]")
        rows: list[list[str]] = []
        for row in ws.iter_rows(max_row=min(ws.max_row or 0, 500)):
            cells = ["" if c.value is None else str(c.value) for c in row]
            if any(cell.strip() for cell in cells):
                rows.append(cells)
                parts.append(" | ".join(cells))
        if rows:
            tables.append(TableData(doc_name=f"{path.name}#{ws.title}", rows=rows))
    full_text = "\n".join(parts)[:MAX_CHARS_PER_DOC]
    return ParsedDoc(
        name=path.name, kind="xlsx", full_text=full_text, tables=tables,
        fillable_template=True, cui_markings=scan_cui(full_text),
    )


def detect_sections(text: str) -> list[DocSection]:
    """UCF A–M and attachment/SOW header detection (regex first; the
    classification stage escalates ambiguity to an LLM)."""
    sections: list[DocSection] = []
    matches = list(SECTION_HEADER_RE.finditer(text))
    for i, m in enumerate(matches):
        header = " ".join(m.group(0).split())
        ucf_letter = m.group(2)
        section_id = ucf_letter.upper() if ucf_letter else header[:24]
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        sections.append(
            DocSection(section_id=section_id, title=header, text=text[start:end][:100_000])
        )
    return sections


def scan_cui(text: str) -> list[str]:
    found = []
    upper = text[:200_000].upper()
    for pattern in CUI_PATTERNS:
        m = re.search(pattern, upper)
        if m:
            found.append(m.group(0))
    return sorted(set(found))
