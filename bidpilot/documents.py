"""Attachment text extraction (PDF, DOCX, plain text)."""

from __future__ import annotations

from pathlib import Path

from .models import AttachmentInfo

# Cap per-attachment text so one enormous PDF can't swamp the context window.
MAX_CHARS_PER_ATTACHMENT = 400_000


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix in (".docx", ".dotx"):
        text = _extract_docx(path)
    elif suffix in (".txt", ".md", ".csv", ".json", ".xml", ".htm", ".html"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        text = ""
    return text[:MAX_CHARS_PER_ATTACHMENT]


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            pages.append(f"[page {i + 1}]\n{page_text}")
        return "\n\n".join(pages)
    except Exception:
        return ""


def _extract_docx(path: Path) -> str:
    import docx

    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_attachments(attachments: list[AttachmentInfo]) -> dict[str, str]:
    """Extract text from every attachment; annotates each AttachmentInfo in place.

    Returns {attachment name: extracted text} for attachments with usable text.
    """
    texts: dict[str, str] = {}
    for att in attachments:
        text = extract_text(att.local_path)
        att.extracted = bool(text.strip())
        att.char_count = len(text)
        if att.extracted:
            texts[att.name] = text
    return texts


def build_corpus(description: str, attachment_texts: dict[str, str]) -> str:
    """Assemble the full solicitation corpus the agents read."""
    parts = ["=== SAM.GOV NOTICE DESCRIPTION ===", description or "(no description text)"]
    for name, text in attachment_texts.items():
        parts.append(f"\n=== ATTACHMENT: {name} ===")
        parts.append(text)
    return "\n\n".join(parts)
