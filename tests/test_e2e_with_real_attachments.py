"""Integration test: the full pipeline over REAL attachment files —
a DOCX solicitation, a government XLSX pricing template, and a fillable PDF
form — exercising docproc parsing, template detection + machine fill,
fillable-form prefill, retrieval-backed writers, rendering, and export."""

from pathlib import Path

from rich.console import Console

from bidpilot.audit import AuditLog
from bidpilot.kb.store import load_kb
from bidpilot.models import AttachmentRecord, NoticeMetadata
from bidpilot.orchestrator import RunContext, new_run, run
from bidpilot.pricing.models import CellWrite, TemplateFillProposal
from bidpilot.state import Stage

from test_orchestrator_e2e import EXAMPLE_KB, FakeRouter, FakeSam

NOTICE = "e" * 32


# ---------------------------------------------------------------------------
# Fixture attachments (built with the same libs the parsers use)
# ---------------------------------------------------------------------------


def _make_solicitation_docx(path: Path):
    import docx

    d = docx.Document()
    d.add_paragraph("SECTION C - DESCRIPTION / SPECIFICATIONS")
    d.add_paragraph(
        "The contractor shall provide IT help desk services handling incident "
        "tickets, and shall maintain the agency's web applications."
    )
    d.add_paragraph("SECTION L - INSTRUCTIONS TO OFFERORS")
    d.add_paragraph(
        "The offeror shall submit Volume I Technical, limited to 20 pages. "
        "Proposals shall be emailed to co@agency.gov no later than "
        "2026-08-15 14:00 Eastern Time. Offerors shall complete the pricing "
        "template in Attachment 2 and the SF-form in Attachment 3."
    )
    d.add_paragraph("SECTION M - EVALUATION FACTORS")
    d.add_paragraph("Technical approach will be evaluated for soundness and specificity.")
    d.save(str(path))


def _make_pricing_template_xlsx(path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws["A1"], ws["B1"], ws["C1"] = "CLIN", "Description", "Extended Price"
    ws["A2"], ws["B2"] = "0001", "Help Desk Services (Base Year)"
    ws["C3"] = "=SUM(C2:C2)"
    wb.save(str(path))


def _make_fillable_form_pdf(path: Path):
    from test_pdf_forms_and_dashboard import _make_fillable_pdf

    _make_fillable_pdf(path)


class AttachmentSam(FakeSam):
    """FakeSam that serves the three real fixture files."""

    def __init__(self, fixtures_dir: Path):
        self.fixtures_dir = fixtures_dir

    def notice_metadata(self, notice_id):
        meta = super().notice_metadata(notice_id)
        return NoticeMetadata(**{**meta.model_dump(), "notice_id": notice_id})

    def download_attachments(self, notice_id, dest_dir):
        dest_dir = Path(dest_dir)
        dest_dir.mkdir(parents=True, exist_ok=True)
        records = []
        for name, maker in (
            ("Attachment_1_Solicitation.docx", _make_solicitation_docx),
            ("Attachment_2_Pricing_Template.xlsx", _make_pricing_template_xlsx),
            ("Attachment_3_SF_Form.pdf", _make_fillable_form_pdf),
        ):
            file_path = dest_dir / name
            if not file_path.exists():
                maker(file_path)
            records.append(AttachmentRecord(
                name=name, local_path=str(file_path), source_notice_id=notice_id,
            ))
        return records


class AttachmentRouter(FakeRouter):
    """Extends the canned router for the template-fill and form-fill schemas."""

    def structured(self, tier, **kw):
        name = kw["output_type"].__name__
        if name == "TemplateFillProposal":
            return TemplateFillProposal(writes=[
                CellWrite(sheet="Pricing", cell="C2", value="18533.00", note="help desk base year"),
            ])
        if name == "FieldMapping":
            return kw["output_type"](fields={"CompanyName": "Example Federal Solutions LLC"})
        if name == "FormsPackage":
            pkg = super().structured(tier, **kw)
            pkg.forms[0].source_file = "Attachment_3_SF_Form.pdf"
            return pkg
        return super().structured(tier, **kw)


def test_full_pipeline_with_real_attachments(tmp_path):
    state, checkpoints = new_run(NOTICE, tmp_path)
    ctx = RunContext(
        state=state,
        router=AttachmentRouter(),
        sam=AttachmentSam(tmp_path / "fixtures"),
        kb=load_kb(str(EXAMPLE_KB)),
        checkpoints=checkpoints,
        audit=AuditLog(Path(state.run_dir) / "audit.jsonl"),
        console=Console(quiet=True),
        confirm=lambda q: True,
        actor="test",
    )
    result = run(ctx)

    assert result.halted_reason is None
    assert all(result.is_done(s) for s in Stage)
    run_dir = Path(result.run_dir)

    # Docproc parsed the real files: DOCX sections + XLSX marked fillable.
    docx_doc = next(d for d in result.doc_tree.docs if d.name.endswith(".docx"))
    assert {s.section_id for s in docx_doc.sections} >= {"C", "L", "M"}
    xlsx_doc = next(d for d in result.doc_tree.docs if d.name.endswith(".xlsx"))
    assert xlsx_doc.fillable_template
    assert xlsx_doc.tables and any("CLIN" in cell for row in xlsx_doc.tables[0].rows for cell in row)

    # Template detected and machine-filled on a copy; formulas preserved.
    assert result.pricing.structure.government_template_file == "Attachment_2_Pricing_Template.xlsx"
    filled = run_dir / "pricing" / "FILLED_Attachment_2_Pricing_Template.xlsx"
    assert filled.exists()
    import openpyxl

    wb = openpyxl.load_workbook(str(filled))
    assert wb["Pricing"]["C2"].value == 18533.00
    assert wb["Pricing"]["C3"].value == "=SUM(C2:C2)"
    original = run_dir / "attachments" / "Attachment_2_Pricing_Template.xlsx"
    assert openpyxl.load_workbook(str(original))["Pricing"]["C2"].value is None  # untouched

    # Fillable PDF prefilled with admin data only.
    prefilled = run_dir / "forms" / "PREFILLED_Attachment_3_SF_Form.pdf"
    assert prefilled.exists()
    from pypdf import PdfReader

    fields = PdfReader(str(prefilled)).get_fields()
    assert fields["CompanyName"].get("/V") == "Example Federal Solutions LLC"
    assert result.forms.forms[0].filled_file == str(prefilled)

    # Rendered volumes + export bundle.
    assert result.rendered_volumes and Path(result.rendered_volumes[0].docx_path).exists()
    assert result.export_path and Path(result.export_path).exists()

    # The filled template ships in the ZIP; raw attachments do not.
    import zipfile

    names = zipfile.ZipFile(result.export_path).namelist()
    assert "pricing/FILLED_Attachment_2_Pricing_Template.xlsx" in names
    assert not any(n.startswith("attachments/") for n in names)
