from pathlib import Path

from bidpilot.models import FormatConstraints, QASeverity, SectionDraft
from bidpilot.pricing.models import CellWrite, TemplateFillProposal
from bidpilot.pricing.workbook import apply_fill, dump_grid
from bidpilot.rendering import (
    render_all_volumes,
    verify_rendered,
    volume_filename,
    _parse_font,
    _parse_margin_inches,
    _parse_pt,
)


def _draft(md="## Approach\n\n- bullet one\n\n**Bold** paragraph text here.", volume="Volume I - Technical"):
    return SectionDraft(
        section_id="TECH-1", volume=volume, title="Approach",
        markdown=md, word_count=len(md.split()),
    )


def test_render_docx_produced(tmp_path):
    constraints = FormatConstraints(font="12-point Times New Roman", margins='1" margins all around')
    rendered = render_all_volumes([_draft()], constraints, tmp_path, "W9123-26-R-0001", "Example LLC")
    assert len(rendered) == 1
    rv = rendered[0]
    assert Path(rv.docx_path).exists()
    assert Path(rv.docx_path).name.startswith("W9123-26-R-0001_Volume_I_Technical")
    assert rv.word_count > 0
    # docx is a real zip container
    assert Path(rv.docx_path).read_bytes()[:2] == b"PK"


def test_render_multiple_volumes(tmp_path):
    rendered = render_all_volumes(
        [_draft(), _draft(volume="Volume II - Past Performance")],
        FormatConstraints(), tmp_path, None, None,
    )
    assert len(rendered) == 2


def test_verify_rendered_exact_page_violation():
    from bidpilot.models import RenderedVolumeInfo

    rv = RenderedVolumeInfo(volume="Volume I", docx_path="x.docx", page_count=25, estimated_pages=25.0)
    findings = verify_rendered([rv], FormatConstraints(page_limits={"Volume I": 20}))
    assert findings[0].severity == QASeverity.HARD
    assert "exact count" in findings[0].description


def test_verify_rendered_estimate_violation():
    from bidpilot.models import RenderedVolumeInfo

    rv = RenderedVolumeInfo(volume="Volume I", docx_path="x.docx", estimated_pages=30.0, word_count=13500)
    findings = verify_rendered([rv], FormatConstraints(page_limits={"Volume I": 20}))
    assert findings[0].severity == QASeverity.HARD
    assert "estimated" in findings[0].description


def test_constraint_parsers():
    assert _parse_font("Font shall be Times New Roman, 12 point") == "Times New Roman"
    assert _parse_pt("12 pt minimum") == 12.0
    assert _parse_margin_inches('one-inch (1") margins') == 1.0
    assert _parse_font(None) is None


def test_volume_filename_sanitized():
    name = volume_filename("Volume I — Technical / Management", "W9123-26-R-0001", "Ex & Co LLC")
    assert " " not in name and "/" not in name
    assert name.endswith(".docx")


def _make_template(path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws["A1"] = "Labor Category"
    ws["B1"] = "Extended Price"
    ws["A2"] = "Software Engineer"
    ws["B3"] = "=SUM(B2:B2)"  # formula cell — must never be overwritten
    wb.save(str(path))


def test_workbook_dump_and_apply_fill(tmp_path):
    template = tmp_path / "gov_pricing.xlsx"
    _make_template(template)
    grid = dump_grid(template)
    assert "Pricing!A1: Labor Category" in grid
    assert "Pricing!B3: =SUM(B2:B2)" in grid

    proposal = TemplateFillProposal(writes=[
        CellWrite(sheet="Pricing", cell="B2", value="18533.00", note="SE extended"),
        CellWrite(sheet="Pricing", cell="B3", value="999", note="attempt to hit formula"),
        CellWrite(sheet="Nope", cell="A1", value="1", note="bad sheet"),
    ])
    dest = tmp_path / "FILLED_gov_pricing.xlsx"
    filled, skipped = apply_fill(template, proposal, dest)
    assert filled == dest and dest.exists()
    assert any("formula cell" in s for s in skipped)
    assert any("sheet not found" in s for s in skipped)

    import openpyxl

    wb = openpyxl.load_workbook(str(dest))
    assert wb["Pricing"]["B2"].value == 18533.00
    assert wb["Pricing"]["B3"].value == "=SUM(B2:B2)"  # formula preserved


def test_workbook_unfillable_returns_none(tmp_path):
    template = tmp_path / "gov.xlsx"
    _make_template(template)
    proposal = TemplateFillProposal(unfillable_reason="merged cells everywhere")
    filled, notes = apply_fill(template, proposal, tmp_path / "out.xlsx")
    assert filled is None
    assert "merged cells" in notes[0]
